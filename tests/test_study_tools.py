"""
tests/test_study_tools.py  —  ClassPilot Study Assistant (5 tools)
Pure unittest, no network. All Google APIs fully mocked.
"""

import io, os, sys, types, unittest
from unittest.mock import MagicMock, patch, call

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

# ── stub external deps ────────────────────────────────────────────────────────
import types as _types, sys as _sys
_aps  = _types.ModuleType("apscheduler");            _aps.__path__  = []; _sys.modules["apscheduler"]            = _aps
_apss = _types.ModuleType("apscheduler.schedulers"); _apss.__path__ = []; _sys.modules["apscheduler.schedulers"] = _apss; _aps.schedulers = _apss
del _aps, _apss, _types, _sys

def _stub(*names):
    for name in names:
        parts = name.split(".")
        for i in range(1, len(parts) + 1):
            key = ".".join(parts[:i])
            if key not in sys.modules:
                m = types.ModuleType(key)
                m.__path__ = []
                sys.modules[key] = m
                if i > 1:
                    setattr(sys.modules[".".join(parts[:i - 1])], parts[i - 1], m)

# Force a clean re-import of classroom_suite_mcp.auth and the classpilot
# modules that patch it: another test file in this suite (test_feature4_5.py)
# imports the REAL vendor `classroom_suite_mcp` package (to exercise
# classpilot.classroom_client's scope-patching against real code), and all
# test files in a pytest run share one process-wide sys.modules cache. If
# that real import happened first, `classpilot.classroom_client`'s one-time,
# import-time SCOPES-patching code has already run and is cached — it will
# NOT re-run against the stub `SCOPES = []` this file sets up below, and
# TestScopeExtension would then be checking a leftover empty list.
#
# Only `auth` (and the two classpilot modules that import it) are dropped —
# deliberately NOT `classroom_suite_mcp` itself or its `.classroom` /
# `.drive` / `.docs` submodules. Those are left exactly as any earlier test
# file cached them (real vendor code from test_feature4_5.py, or absent if
# this file runs first) so test_feature4_5.py's later assertions against the
# real `classroom_suite_mcp.classroom` module (e.g. attach_submission_files)
# keep working regardless of run order.
for _mod in ("classpilot.classroom_client", "classpilot.study_client", "classroom_suite_mcp.auth"):
    sys.modules.pop(_mod, None)

_stub(
    "anthropic", "dotenv", "fastmcp", "pydantic",
    "google.auth.transport.requests", "google.oauth2.credentials",
    "google_auth_oauthlib.flow",
    "googleapiclient.discovery", "googleapiclient.http", "googleapiclient.errors",
    "apscheduler.schedulers.blocking", "apscheduler.schedulers.background",
    "apscheduler.executors.pool",
    "classroom_suite_mcp",
    "classroom_suite_mcp.auth",
    "classroom_suite_mcp.classroom",
    "classroom_suite_mcp.drive",
    "classroom_suite_mcp.docs",
)
sys.modules["dotenv"].load_dotenv = lambda *a, **k: None

# pydantic stub
class _BM:
    def __init_subclass__(cls, **k): pass
    def __init__(self, **k):
        for key, val in k.items(): setattr(self, key, val)
class _Field:
    def __call__(self, *a, **k): return None
sys.modules["pydantic"].BaseModel = _BM
sys.modules["pydantic"].Field = _Field()

sys.modules["fastmcp"].FastMCP = type("FastMCP", (), {
    "__init__": lambda s, *a, **k: None,
    "tool": lambda s, *a, **k: (lambda f: f),
})
sys.modules["google.auth.transport.requests"].Request = object
sys.modules["google.oauth2.credentials"].Credentials = object
sys.modules["google_auth_oauthlib.flow"].InstalledAppFlow = object
sys.modules["googleapiclient.discovery"].build = lambda *a, **k: MagicMock()
sys.modules["googleapiclient.discovery"].Resource = object
sys.modules["googleapiclient.http"].MediaIoBaseUpload = object
sys.modules["googleapiclient.http"].MediaIoBaseDownload = MagicMock
sys.modules["googleapiclient.errors"].HttpError = Exception
sys.modules["apscheduler.schedulers.blocking"].BlockingScheduler = type("BS", (), {"__init__": lambda s, **k: None})
sys.modules["apscheduler.schedulers.background"].BackgroundScheduler = type("BGS", (), {
    "__init__": lambda s, **k: None, "add_job": lambda s, *a, **k: None,
    "start": lambda s: None, "shutdown": lambda s, **k: None,
    "remove_job": lambda s, j: None, "running": True,
})
sys.modules["apscheduler.executors.pool"].ThreadPoolExecutor = type("TPE", (), {"__init__": lambda s, **k: None})

# Stub auth functions used by study_client
_mock_classroom_svc = MagicMock()
_mock_drive_svc = MagicMock()


def _make_png(size, color) -> io.BytesIO:
    """Small helper for visual-extraction tests — builds an in-memory PNG."""
    from PIL import Image as PILImage
    im = PILImage.new("RGB", size, color=color)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf
sys.modules["classroom_suite_mcp.auth"].get_classroom_service = lambda: _mock_classroom_svc
sys.modules["classroom_suite_mcp.auth"].get_drive_service     = lambda: _mock_drive_svc
sys.modules["classroom_suite_mcp.auth"].get_docs_service      = lambda: MagicMock()
sys.modules["classroom_suite_mcp.auth"].SCOPES = []
sys.modules["classroom_suite_mcp.auth"].GoogleAuth = MagicMock

# ── imports under test ────────────────────────────────────────────────────────
import classpilot.study_client as sc


def _reset_mocks():
    # return_value=True / side_effect=True: also clear any `.execute.return_value`
    # or `.execute.side_effect` a previous test configured deep in the mock chain
    # (e.g. pagination tests set `.execute.side_effect = [...]`, which — unlike
    # call counts — a plain reset_mock() does NOT clear on its own, and it would
    # otherwise leak into and break the next test that reuses the same call chain).
    _mock_classroom_svc.reset_mock(return_value=True, side_effect=True)
    _mock_drive_svc.reset_mock(return_value=True, side_effect=True)


# ── sample fixtures ───────────────────────────────────────────────────────────

RAW_COURSE = {"id": "c1", "name": "Deep Learning", "section": "A", "courseState": "ACTIVE", "alternateLink": "http://x"}
RAW_TOPIC  = {"topicId": "1", "name": "Module 1 - Intro", "updateTime": "2026-01-01T00:00:00Z"}
RAW_TOPIC2 = {"topicId": "2", "name": "Module 2 - CNN",   "updateTime": "2026-01-02T00:00:00Z"}

RAW_MATERIAL = {
    "id": "m1", "title": "Session 1 - Intro to DL", "description": "First lecture",
    "topicId": "1", "state": "PUBLISHED",
    "materials": [
        {"driveFile": {"driveFile": {"id": "f1", "title": "Session1.pdf", "alternateLink": "http://pdf"}}},
        {"youtubeVideo": {"id": "yt1", "title": "DL Intro Video", "alternateLink": "http://yt"}},
    ],
}
RAW_ASSIGNMENT = {
    "id": "a1", "title": "Lab 1 Assignment", "description": "Submit lab work",
    "topicId": "1", "state": "PUBLISHED",
    "dueDate": {"year": 2026, "month": 12, "day": 1},
    "dueTime": {"hours": 23, "minutes": 59},
    "materials": [],
}


# ═════════════════════════════════════════════════════════════════════════════
class TestFetchCourses(unittest.TestCase):
    def setUp(self): _reset_mocks()

    def test_returns_list_of_courses(self):
        _mock_classroom_svc.courses().list().execute.return_value = {"courses": [RAW_COURSE]}
        result = sc.fetch_courses()
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]["name"], "Deep Learning")

    def test_returns_empty_when_no_courses(self):
        _mock_classroom_svc.courses().list().execute.return_value = {"courses": []}
        self.assertEqual(sc.fetch_courses(), [])

    def test_filters_active_courses(self):
        _mock_classroom_svc.courses().list().execute.return_value = {"courses": [RAW_COURSE]}
        sc.fetch_courses()
        call_kwargs = _mock_classroom_svc.courses().list.call_args[1]
        self.assertIn("ACTIVE", call_kwargs.get("courseStates", []))

    def test_all_fields_present(self):
        _mock_classroom_svc.courses().list().execute.return_value = {"courses": [RAW_COURSE]}
        r = sc.fetch_courses()[0]
        for key in ("id", "name", "section", "state", "link"):
            self.assertIn(key, r)

    def test_pages_through_nextPageToken(self):
        """A truncated first page (with nextPageToken) must not silently drop courses."""
        course2 = dict(RAW_COURSE, id="c2", name="Quantum Computing")
        _mock_classroom_svc.courses().list().execute.side_effect = [
            {"courses": [RAW_COURSE], "nextPageToken": "TOK1"},
            {"courses": [course2]},  # final page: no nextPageToken
        ]
        result = sc.fetch_courses()
        self.assertEqual(len(result), 2)
        self.assertEqual({r["name"] for r in result}, {"Deep Learning", "Quantum Computing"})

    def test_second_page_request_includes_page_token(self):
        course2 = dict(RAW_COURSE, id="c2")
        _mock_classroom_svc.courses().list().execute.side_effect = [
            {"courses": [RAW_COURSE], "nextPageToken": "TOK1"},
            {"courses": [course2]},
        ]
        sc.fetch_courses()
        second_call_kwargs = _mock_classroom_svc.courses().list.call_args_list[-1][1]
        self.assertEqual(second_call_kwargs.get("pageToken"), "TOK1")


class TestFetchTopics(unittest.TestCase):
    def setUp(self): _reset_mocks()

    def test_returns_sorted_topics(self):
        _mock_classroom_svc.courses().topics().list().execute.return_value = {
            "topic": [RAW_TOPIC2, RAW_TOPIC]  # out of order
        }
        result = sc.fetch_topics("c1")
        self.assertEqual(result[0]["name"], "Module 1 - Intro")
        self.assertEqual(result[1]["name"], "Module 2 - CNN")

    def test_returns_empty_when_no_topics(self):
        _mock_classroom_svc.courses().topics().list().execute.return_value = {"topic": []}
        self.assertEqual(sc.fetch_topics("c1"), [])

    def test_all_fields_present(self):
        _mock_classroom_svc.courses().topics().list().execute.return_value = {"topic": [RAW_TOPIC]}
        r = sc.fetch_topics("c1")[0]
        self.assertEqual(r["id"], "1")
        self.assertEqual(r["name"], "Module 1 - Intro")

    def test_pages_through_nextPageToken(self):
        """Modules beyond the first page must still be returned, not truncated."""
        _mock_classroom_svc.courses().topics().list().execute.side_effect = [
            {"topic": [RAW_TOPIC], "nextPageToken": "TOK1"},
            {"topic": [RAW_TOPIC2]},
        ]
        result = sc.fetch_topics("c1")
        self.assertEqual(len(result), 2)
        self.assertEqual({r["name"] for r in result}, {"Module 1 - Intro", "Module 2 - CNN"})


class TestFetchMaterials(unittest.TestCase):
    def setUp(self): _reset_mocks()

    def test_fetches_course_work_materials(self):
        _mock_classroom_svc.courses().courseWorkMaterials().list().execute.return_value = {
            "courseWorkMaterial": [RAW_MATERIAL]
        }
        result = sc.fetch_course_work_materials("c1")
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]["kind"], "MATERIAL")

    def test_filters_by_topic_id(self):
        mat2 = dict(RAW_MATERIAL); mat2["topicId"] = "99"
        _mock_classroom_svc.courses().courseWorkMaterials().list().execute.return_value = {
            "courseWorkMaterial": [RAW_MATERIAL, mat2]
        }
        result = sc.fetch_course_work_materials("c1", topic_id="1")
        self.assertEqual(len(result), 1)

    def test_attachments_normalised(self):
        _mock_classroom_svc.courses().courseWorkMaterials().list().execute.return_value = {
            "courseWorkMaterial": [RAW_MATERIAL]
        }
        result = sc.fetch_course_work_materials("c1")
        atts = result[0]["attachments"]
        types = [a["type"] for a in atts]
        self.assertIn("drive",   types)
        self.assertIn("youtube", types)

    def test_fetches_assignments_with_assignment_kind(self):
        _mock_classroom_svc.courses().courseWork().list().execute.return_value = {
            "courseWork": [RAW_ASSIGNMENT]
        }
        result = sc.fetch_assignments("c1")
        self.assertEqual(result[0]["kind"], "ASSIGNMENT")

    def test_materials_page_through_nextPageToken_before_topic_filter(self):
        """
        Pagination must happen across the FULL result set before the
        topic_id filter is applied, otherwise a later page containing
        matching-topic items would be silently dropped.
        """
        mat2 = dict(RAW_MATERIAL, id="m2")  # same topicId "1" as RAW_MATERIAL
        _mock_classroom_svc.courses().courseWorkMaterials().list().execute.side_effect = [
            {"courseWorkMaterial": [RAW_MATERIAL], "nextPageToken": "TOK1"},
            {"courseWorkMaterial": [mat2]},
        ]
        result = sc.fetch_course_work_materials("c1", topic_id="1")
        self.assertEqual(len(result), 2)

    def test_assignments_page_through_nextPageToken(self):
        assignment2 = dict(RAW_ASSIGNMENT, id="a2")
        _mock_classroom_svc.courses().courseWork().list().execute.side_effect = [
            {"courseWork": [RAW_ASSIGNMENT], "nextPageToken": "TOK1"},
            {"courseWork": [assignment2]},
        ]
        result = sc.fetch_assignments("c1")
        self.assertEqual(len(result), 2)


class TestFetchAnnouncementsPagination(unittest.TestCase):
    """Announcements pagination — fetch_announcements has its own try/except wrapper."""

    def setUp(self): _reset_mocks()

    RAW_ANNOUNCEMENT = {"id": "an1", "text": "Midterm moved to next week", "state": "PUBLISHED"}

    def test_returns_announcements(self):
        _mock_classroom_svc.courses().announcements().list().execute.return_value = {
            "announcement": [self.RAW_ANNOUNCEMENT]
        }
        result = sc.fetch_announcements("c1")
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]["kind"], "ANNOUNCEMENT")

    def test_pages_through_nextPageToken(self):
        ann2 = dict(self.RAW_ANNOUNCEMENT, id="an2")
        _mock_classroom_svc.courses().announcements().list().execute.side_effect = [
            {"announcement": [self.RAW_ANNOUNCEMENT], "nextPageToken": "TOK1"},
            {"announcement": [ann2]},
        ]
        result = sc.fetch_announcements("c1")
        self.assertEqual(len(result), 2)

    def test_failure_returns_empty_list_not_exception(self):
        _mock_classroom_svc.courses().announcements().list().execute.side_effect = Exception("no scope")
        result = sc.fetch_announcements("c1")
        self.assertEqual(result, [])


class TestPaginateHelper(unittest.TestCase):
    """Direct tests of the generic _paginate() helper used by every fetch_* function."""

    def test_single_page_no_token(self):
        list_fn = MagicMock()
        list_fn.return_value.execute.return_value = {"items": [{"id": 1}, {"id": 2}]}
        result = sc._paginate(list_fn, "items", {"courseId": "c1"})
        self.assertEqual(result, [{"id": 1}, {"id": 2}])
        list_fn.assert_called_once_with(courseId="c1")

    def test_follows_multiple_pages(self):
        list_fn = MagicMock()
        list_fn.return_value.execute.side_effect = [
            {"items": [{"id": 1}], "nextPageToken": "A"},
            {"items": [{"id": 2}], "nextPageToken": "B"},
            {"items": [{"id": 3}]},
        ]
        result = sc._paginate(list_fn, "items", {"courseId": "c1"})
        self.assertEqual([i["id"] for i in result], [1, 2, 3])
        self.assertEqual(list_fn.call_count, 3)

    def test_passes_page_token_on_subsequent_calls(self):
        list_fn = MagicMock()
        list_fn.return_value.execute.side_effect = [
            {"items": [{"id": 1}], "nextPageToken": "TOK_A"},
            {"items": [{"id": 2}]},
        ]
        sc._paginate(list_fn, "items", {"courseId": "c1"})
        first_kwargs = list_fn.call_args_list[0][1]
        second_kwargs = list_fn.call_args_list[1][1]
        self.assertNotIn("pageToken", first_kwargs)
        self.assertEqual(second_kwargs.get("pageToken"), "TOK_A")

    def test_missing_item_key_treated_as_empty_page(self):
        list_fn = MagicMock()
        list_fn.return_value.execute.return_value = {}  # no "items" key at all
        result = sc._paginate(list_fn, "items", {})
        self.assertEqual(result, [])

    def test_does_not_mutate_caller_params_across_pages(self):
        """Each page's call params must be independent — no leaking pageToken across calls."""
        list_fn = MagicMock()
        list_fn.return_value.execute.side_effect = [
            {"items": [{"id": 1}], "nextPageToken": "A"},
            {"items": [{"id": 2}]},
        ]
        base_params = {"courseId": "c1"}
        sc._paginate(list_fn, "items", base_params)
        self.assertNotIn("pageToken", base_params)


class TestNormaliseMaterial(unittest.TestCase):
    def test_drive_attachment_extracted(self):
        result = sc._normalise_material(RAW_MATERIAL, "MATERIAL")
        drive_att = next((a for a in result["attachments"] if a["type"] == "drive"), None)
        self.assertIsNotNone(drive_att)
        self.assertEqual(drive_att["id"], "f1")

    def test_youtube_attachment_extracted(self):
        result = sc._normalise_material(RAW_MATERIAL, "MATERIAL")
        yt_att = next((a for a in result["attachments"] if a["type"] == "youtube"), None)
        self.assertIsNotNone(yt_att)

    def test_kind_preserved(self):
        r = sc._normalise_material(RAW_ASSIGNMENT, "ASSIGNMENT")
        self.assertEqual(r["kind"], "ASSIGNMENT")

    def test_due_date_included_for_assignments(self):
        r = sc._normalise_material(RAW_ASSIGNMENT, "ASSIGNMENT")
        self.assertEqual(r["due_date"], {"year": 2026, "month": 12, "day": 1})


class TestReadDriveFile(unittest.TestCase):
    def setUp(self): _reset_mocks()

    def _setup_meta(self, mime, name="test.file", link="http://file"):
        _mock_drive_svc.files().get().execute.return_value = {
            "id": "f1", "name": name, "mimeType": mime, "webViewLink": link, "size": "1000"
        }

    def test_google_doc_exported_as_text(self):
        self._setup_meta("application/vnd.google-apps.document", "Lecture Notes")
        import io
        buf = io.BytesIO(b"Lecture content here")
        mock_dl = MagicMock()
        mock_dl.next_chunk.return_value = (None, True)
        with patch("classpilot.study_client.MediaIoBaseDownload", return_value=mock_dl):
            with patch("io.BytesIO", return_value=buf):
                _mock_drive_svc.files().export_media.return_value = MagicMock()
                result = sc.read_drive_file("f1")
        self.assertTrue(result["readable"])
        self.assertEqual(result["mime_type"], "application/vnd.google-apps.document")

    def test_pdf_text_extracted(self):
        self._setup_meta("application/pdf", "Slides.pdf", "http://pdf")
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        raw = io.BytesIO()
        writer.write(raw)
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        # A blank page has no extractable text, so this should be readable=False
        # with a graceful "no text found" note rather than an error/crash.
        self.assertFalse(result["readable"])
        self.assertEqual(result["content"], "")
        self.assertIn("http://pdf", result["link"])
        self.assertIn("no extractable text", result["note"])

    def test_pdf_with_real_text_extracted(self):
        self._setup_meta("application/pdf", "Notes.pdf", "http://pdf")
        from pypdf import PdfWriter
        try:
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest("reportlab not available to build a text PDF fixture")

        raw = io.BytesIO()
        c = canvas.Canvas(raw)
        c.drawString(100, 700, "Neural Networks Lecture Notes")
        c.save()
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertTrue(result["readable"])
        self.assertIn("Neural Networks Lecture Notes", result["content"])
        self.assertEqual(result["note"], "")

    def test_pdf_extraction_failure_falls_back_to_url_only(self):
        self._setup_meta("application/pdf", "Corrupt.pdf", "http://pdf")

        def _fake_download(buf, request):
            buf.write(b"not a real pdf file")
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertFalse(result["readable"])
        self.assertEqual(result["content"], "")
        self.assertIn("http://pdf", result["link"])

    def test_pdf_empty_file_falls_back_to_url_only(self):
        self._setup_meta("application/pdf", "Empty.pdf", "http://pdf")

        def _fake_download(buf, request):
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertFalse(result["readable"])
        self.assertEqual(result["content"], "")

    def test_docx_text_extracted(self):
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Essay.docx", "http://docx"
        )
        from docx import Document as _RealDocument

        doc = _RealDocument()
        doc.add_heading("Assignment 1", level=1)
        doc.add_paragraph("This essay discusses backpropagation.")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Term"
        table.rows[0].cells[1].text = "Definition"
        raw = io.BytesIO()
        doc.save(raw)
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertTrue(result["readable"])
        self.assertIn("Assignment 1", result["content"])
        self.assertIn("backpropagation", result["content"])
        self.assertIn("Term | Definition", result["content"])
        self.assertEqual(result["note"], "")

    def test_docx_extraction_failure_falls_back_to_url_only(self):
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Corrupt.docx", "http://docx"
        )

        def _fake_download(buf, request):
            buf.write(b"not a real docx file")
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertFalse(result["readable"])
        self.assertEqual(result["content"], "")
        self.assertIn("http://docx", result["link"])

    def test_docx_empty_document_returns_no_text_note(self):
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Blank.docx", "http://docx"
        )
        from docx import Document as _RealDocument

        doc = _RealDocument()
        raw = io.BytesIO()
        doc.save(raw)
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertFalse(result["readable"])
        self.assertEqual(result["content"], "")
        self.assertIn("no extractable text", result["note"])

    def test_legacy_doc_returns_url_only(self):
        self._setup_meta("application/msword", "OldEssay.doc")
        result = sc.read_drive_file("f1")
        self.assertFalse(result["readable"])
        self.assertIn("legacy .doc", result["note"])

    def test_pptx_text_extracted(self):
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Lecture.pptx", "http://pptx"
        )
        from pptx import Presentation as _RealPresentation

        prs = _RealPresentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Neural Networks"
        slide.placeholders[1].text_frame.text = "Intro to CNNs"
        raw = io.BytesIO()
        prs.save(raw)
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertTrue(result["readable"])
        self.assertIn("Neural Networks", result["content"])
        self.assertIn("Intro to CNNs", result["content"])
        self.assertEqual(result["note"], "")

    def test_pptx_extraction_failure_falls_back_to_url_only(self):
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Corrupt.pptx", "http://pptx"
        )
        def _fake_download(buf, request):
            buf.write(b"not a real pptx file")
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertFalse(result["readable"])
        self.assertEqual(result["content"], "")
        self.assertIn("http://pptx", result["link"])

    def test_legacy_ppt_returns_url_only(self):
        self._setup_meta("application/vnd.ms-powerpoint", "OldLecture.ppt")
        result = sc.read_drive_file("f1")
        self.assertFalse(result["readable"])
        self.assertIn("legacy .ppt", result["note"])

    def test_url_only_result_has_link(self):
        self._setup_meta("application/vnd.ms-powerpoint", "file.ppt", "http://drive/file")
        result = sc.read_drive_file("f1")
        self.assertEqual(result["link"], "http://drive/file")

    def test_result_always_has_required_keys(self):
        self._setup_meta("application/vnd.ms-powerpoint")
        result = sc.read_drive_file("f1")
        for key in ("readable", "content", "mime_type", "name", "link", "note", "visuals"):
            self.assertIn(key, result)

    def test_google_doc_has_empty_visuals_list(self):
        """Visual extraction only applies to PPTX/PDF/DOCX/images — every
        other branch must still return visuals=[] (additive, never absent)."""
        self._setup_meta("application/vnd.google-apps.document", "Notes")
        buf = io.BytesIO(b"Some text")
        mock_dl = MagicMock()
        mock_dl.next_chunk.return_value = (None, True)
        with patch("classpilot.study_client.MediaIoBaseDownload", return_value=mock_dl):
            with patch("io.BytesIO", return_value=buf):
                _mock_drive_svc.files().export_media.return_value = MagicMock()
                result = sc.read_drive_file("f1")
        self.assertEqual(result["visuals"], [])

    def test_pptx_diagram_only_slide_surfaced_as_visual(self):
        """The exact motivating scenario: an SVM.pptx-style diagram-only
        slide must appear under 'visuals' while text extraction from the
        rest of the deck is completely unaffected."""
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "SVM.pptx", "http://pptx"
        )
        from pptx import Presentation as _RealPresentation
        from pptx.util import Inches

        prs = _RealPresentation()
        text_slide = prs.slides.add_slide(prs.slide_layouts[1])
        text_slide.shapes.title.text = "Support Vector Machines"
        text_slide.placeholders[1].text_frame.text = "Margin maximization"

        blank = prs.slide_layouts[6]
        diagram_slide = prs.slides.add_slide(blank)
        diagram_slide.shapes.add_picture(
            _make_png((800, 600), (200, 50, 50)),
            Inches(1), Inches(1), width=Inches(6), height=Inches(4.5),
        )

        raw = io.BytesIO()
        prs.save(raw)
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        # Text extraction: completely unchanged behavior.
        self.assertTrue(result["readable"])
        self.assertIn("Support Vector Machines", result["content"])
        self.assertIn("Margin maximization", result["content"])

        # New: the diagram-only slide is surfaced as a visual.
        self.assertEqual(len(result["visuals"]), 1)
        v = result["visuals"][0]
        self.assertEqual(v["location"], "Slide 2")
        self.assertTrue(v["diagram_only"])
        self.assertEqual(v["mime_type"], "image/png")
        self.assertGreater(len(v["data"]), 0)

    def test_pdf_diagram_page_surfaced_as_visual(self):
        self._setup_meta("application/pdf", "Diagrams.pdf", "http://pdf")
        try:
            from reportlab.pdfgen import canvas
            from reportlab.lib.utils import ImageReader
        except ImportError:
            self.skipTest("reportlab not available to build a PDF fixture")

        raw = io.BytesIO()
        c = canvas.Canvas(raw, pagesize=(600, 800))
        c.drawString(50, 750, "Page 1 has real text content")
        c.showPage()
        c.drawImage(ImageReader(_make_png((500, 400), (10, 200, 10))), 50, 200, width=500, height=400)
        c.showPage()
        c.save()
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertTrue(result["readable"])  # page 1 text untouched
        self.assertEqual(len(result["visuals"]), 1)
        self.assertEqual(result["visuals"][0]["location"], "Page 2")

    def test_docx_embedded_figure_surfaced_as_visual(self):
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Notes.docx", "http://docx"
        )
        from docx import Document as _RealDocument
        from docx.shared import Inches as _DocxInches

        doc = _RealDocument()
        doc.add_paragraph("Backpropagation explained")
        doc.add_picture(_make_png((900, 700), (50, 50, 200)), width=_DocxInches(5))
        raw = io.BytesIO()
        doc.save(raw)
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertTrue(result["readable"])
        self.assertIn("Backpropagation explained", result["content"])
        self.assertEqual(len(result["visuals"]), 1)

    def test_visual_extraction_failure_does_not_break_text_extraction(self):
        """If image extraction blows up for any reason, text extraction
        (already successfully computed) must still be returned intact —
        visual extraction is a strictly additive, fault-isolated feature."""
        self._setup_meta(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Lecture.pptx", "http://pptx"
        )
        from pptx import Presentation as _RealPresentation

        prs = _RealPresentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Neural Networks"
        raw = io.BytesIO()
        prs.save(raw)
        raw_bytes = raw.getvalue()

        def _fake_download(buf, request):
            buf.write(raw_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            with patch("classpilot.study_client.extract_pptx_visuals", side_effect=RuntimeError("boom")):
                _mock_drive_svc.files().get_media.return_value = MagicMock()
                result = sc.read_drive_file("f1")

        self.assertTrue(result["readable"])
        self.assertIn("Neural Networks", result["content"])
        self.assertEqual(result["visuals"], [])  # failed extraction -> empty, not an exception

    def test_png_attachment_returned_as_visual(self):
        self._setup_meta("image/png", "diagram.png", "http://drive/img")
        img_bytes = _make_png((300, 300), (10, 20, 30)).getvalue()

        def _fake_download(buf, request):
            buf.write(img_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertFalse(result["readable"])  # no text — it's an image
        self.assertEqual(len(result["visuals"]), 1)
        self.assertEqual(result["visuals"][0]["data"], img_bytes)
        self.assertEqual(result["visuals"][0]["mime_type"], "image/png")

    def test_jpeg_attachment_returned_as_visual(self):
        self._setup_meta("image/jpeg", "photo.jpg", "http://drive/img2")
        from PIL import Image as PILImage
        jbuf = io.BytesIO()
        PILImage.new("RGB", (200, 200), (5, 5, 5)).save(jbuf, format="JPEG")
        img_bytes = jbuf.getvalue()

        def _fake_download(buf, request):
            buf.write(img_bytes)
            dl = MagicMock()
            dl.next_chunk.return_value = (None, True)
            return dl

        with patch("classpilot.study_client.MediaIoBaseDownload", side_effect=_fake_download):
            _mock_drive_svc.files().get_media.return_value = MagicMock()
            result = sc.read_drive_file("f1")

        self.assertEqual(len(result["visuals"]), 1)
        self.assertEqual(result["visuals"][0]["mime_type"], "image/jpeg")

    def test_gif_still_url_only_not_treated_as_visual(self):
        """Only PNG/JPEG get the direct-visual pipeline per spec — GIF/WEBP
        stay in the url-only fallback."""
        self._setup_meta("image/gif", "animation.gif")
        result = sc.read_drive_file("f1")
        self.assertFalse(result["readable"])
        self.assertEqual(result["visuals"], [])
        self.assertIn("image file", result["note"])


class TestSearchAll(unittest.TestCase):
    def setUp(self):
        _reset_mocks()
        # Clear any side_effect set by previous tests (e.g. test_handles_course_list_failure_gracefully)
        _mock_classroom_svc.courses().list().execute.side_effect = None
        _mock_classroom_svc.courses().list().execute.return_value = {}

    def _setup_classroom(self):
        _mock_classroom_svc.courses().list().execute.return_value = {"courses": [RAW_COURSE]}
        _mock_classroom_svc.courses().topics().list().execute.return_value = {"topic": [RAW_TOPIC]}
        _mock_classroom_svc.courses().courseWorkMaterials().list().execute.return_value = {
            "courseWorkMaterial": [RAW_MATERIAL]
        }
        _mock_classroom_svc.courses().courseWork().list().execute.return_value = {"courseWork": []}

    def test_finds_course_by_name(self):
        self._setup_classroom()
        results = sc.search_all("deep learning")
        types = [r["type"] for r in results]
        self.assertIn("course", types)

    def test_finds_module_by_name(self):
        self._setup_classroom()
        results = sc.search_all("module 1")
        types = [r["type"] for r in results]
        self.assertIn("module", types)

    def test_finds_material_by_title(self):
        self._setup_classroom()
        results = sc.search_all("session 1")
        types = [r["type"] for r in results]
        self.assertIn("material", types)

    def test_returns_empty_for_no_match(self):
        self._setup_classroom()
        results = sc.search_all("quantum cryptography thesis xyz")
        self.assertEqual(results, [])

    def test_result_has_context_field(self):
        self._setup_classroom()
        results = sc.search_all("session 1")
        material_results = [r for r in results if r["type"] == "material"]
        if material_results:
            self.assertIn("context", material_results[0])
            self.assertIn("Deep Learning", material_results[0]["context"])

    def test_case_insensitive(self):
        self._setup_classroom()
        r1 = sc.search_all("DEEP LEARNING")
        r2 = sc.search_all("deep learning")
        self.assertEqual(len(r1), len(r2))

    def test_handles_course_list_failure_gracefully(self):
        _mock_classroom_svc.courses().list().execute.side_effect = Exception("API down")
        results = sc.search_all("anything")
        self.assertEqual(results, [])


class TestScopeExtension(unittest.TestCase):
    """Verify classroom_client.py correctly extends vendor SCOPES."""

    def test_courseworkmaterials_scope_added(self):
        from classroom_suite_mcp import auth as _auth
        scopes = _auth.SCOPES
        self.assertTrue(
            any("courseworkmaterials" in s for s in scopes),
            f"courseworkmaterials scope missing. Got: {scopes}"
        )

    def test_announcements_scope_added(self):
        from classroom_suite_mcp import auth as _auth
        scopes = _auth.SCOPES
        self.assertTrue(
            any("announcements" in s for s in scopes),
            f"announcements scope missing. Got: {scopes}"
        )


class TestSubmissionToolsRemoved(unittest.TestCase):
    """Ensure deprecated submission tools are no longer exposed via MCP."""

    def test_server_has_no_submit_assignment_tool(self):
        import classpilot.server as srv
        tool_fns = [name for name in dir(srv) if name in
                    ("submit_assignment", "turn_in_assignment", "assignment_status")]
        # They may exist as plain functions but should not be MCP tools
        # Check the MCP instance's tool registry
        if hasattr(srv.mcp, "_tools"):
            for bad_tool in ("submit_assignment", "turn_in_assignment"):
                self.assertNotIn(bad_tool, srv.mcp._tools,
                                 f"{bad_tool} should not be an MCP tool")

    def test_study_tools_registered(self):
        import classpilot.server as srv
        expected_study = {
            "list_classes", "list_modules", "list_materials",
            "get_material", "search_classroom",
        }
        if hasattr(srv.mcp, "_tools"):
            for tool in expected_study:
                self.assertIn(tool, srv.mcp._tools,
                              f"Study tool '{tool}' not registered")


if __name__ == "__main__":
    unittest.main(verbosity=2)
